"""Turn role PDFs into .pptx templates with the partner overlay on slide 1."""

from __future__ import annotations

import io
import json
import sys
from copy import deepcopy
from pathlib import Path

import pypdfium2 as pdfium
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
CATALOG = json.loads((ASSETS / "catalog.json").read_text())
OUT_DIR = ASSETS / "templates"
DONOR = ASSETS / "VC_Conf_MP_deck_template.pptx"
PDF_DIR = Path(
    "/Users/timnakhapetov/Downloads/Telegram Desktop/VC-Conf-2027-all-PDFs-1"
)
PDF_OVERRIDES = {
    "media-partner": Path(
        "/Users/timnakhapetov/Dev/presentations/template/VC-Conf-2027-Media-Partners.pdf"
    ),
    "speaker-day3": Path(
        "/Users/timnakhapetov/Downloads/Telegram Desktop/VC_Conf_2027_Day_3_Speakers_Dave_Mullen.pdf"
    ),
}
# 1920-space. Wipe baked-in speaker photo + name on slide 1.
WIPE_SLIDE1 = {
    "speaker-day3": {
        "photo": (98, 334, 145),
        "name": (250, 348, 340, 100),
    }
}
# New Tim headshot lives in this last-slide JPEG; paste onto other decks
# that reuse the same left-card slot (journalist).
TIM_SLIDE = Path(
    "/Users/timnakhapetov/Dev/presentations/template/slides/slide-9.jpg"
)
TIM_BOX = (530, 918, 955, 1343)
PATCH_TIM_ON = {"journalist"}
JPEG_SLIDES = {
    "media-partner": Path("/Users/timnakhapetov/Dev/presentations/template/slides"),
}

SLIDE_W = Inches(20)
SLIDE_H = Inches(11.25)
RENDER_SCALE = 2.0  # 1920pt page → 3840px
OVERLAY_NAMES = {"partner_logo", "partner_label", "partner_name"}


def _blank_layout(prs: Presentation):
    for layout in prs.slide_layouts:
        if layout.name.lower() in {"blank", "пустой слайд"}:
            return layout
    return prs.slide_layouts[-1]


def _overlay_elements(donor_slide):
    return [
        deepcopy(shape._element)
        for shape in donor_slide.shapes
        if shape.name in OVERLAY_NAMES
    ]


def wipe_slide1(im: Image.Image, spec: dict) -> Image.Image:
    s = im.width / 1920
    x, y, d = spec["photo"]
    nx, ny, nw, nh = spec["name"]
    color = im.getpixel((max(0, int(70 * s)), max(0, int(300 * s))))
    draw = ImageDraw.Draw(im)
    pad = 4 * s
    draw.ellipse(
        [x * s - pad, y * s - pad, (x + d) * s + pad, (y + d) * s + pad],
        fill=color,
    )
    x0, y0 = int(nx * s), int(ny * s)
    x1, y1 = int((nx + nw) * s), int((ny + nh) * s)
    strip_y = max(0, y0 - 12)
    strip = im.crop((x0, strip_y, x1, strip_y + 4))
    im.paste(strip.resize((x1 - x0, y1 - y0), Image.Resampling.BILINEAR), (x0, y0))
    return im


def build_one(pdf_path: Path, dest: Path, overlay, template_id: str = "") -> None:
    pdf = pdfium.PdfDocument(str(pdf_path))
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    layout = _blank_layout(prs)
    wipe = WIPE_SLIDE1.get(template_id)

    for i, page in enumerate(pdf):
        pw, _ph = page.get_size()
        scale = 3840 / pw if pw else RENDER_SCALE
        pil = page.render(scale=scale).to_pil().convert("RGB")
        if pil.size != (3840, 2160):
            pil = pil.resize((3840, 2160), Image.Resampling.LANCZOS)
        if i == 0 and wipe:
            pil = wipe_slide1(pil, wipe)
        buf = io.BytesIO()
        pil.save(buf, format="JPEG", quality=92, optimize=True)
        buf.seek(0)
        slide = prs.slides.add_slide(layout)
        slide.shapes.add_picture(buf, Emu(0), Emu(0), width=SLIDE_W, height=SLIDE_H)
        if i == 0 and overlay:
            tree = slide.shapes._spTree
            for el in overlay:
                tree.append(deepcopy(el))

    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dest))
    print(f"  {dest.name}  ({len(pdf)} slides, {dest.stat().st_size // 1024} KB)")


def build_from_jpegs(folder: Path, dest: Path, overlay) -> None:
    files = sorted(folder.glob("slide-*.jpg"), key=lambda p: int(p.stem.split("-")[1]))
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    layout = _blank_layout(prs)
    for i, path in enumerate(files):
        im = Image.open(path).convert("RGB")
        if im.size != (3840, 2160):
            im = im.resize((3840, 2160), Image.Resampling.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=92, optimize=True)
        buf.seek(0)
        slide = prs.slides.add_slide(layout)
        slide.shapes.add_picture(buf, Emu(0), Emu(0), width=SLIDE_W, height=SLIDE_H)
        if i == 0:
            tree = slide.shapes._spTree
            for el in overlay:
                tree.append(deepcopy(el))
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dest))
    print(f"  {dest.name}  ({len(files)} slides from JPEG, {dest.stat().st_size // 1024} KB)")


def _replace_picture(shape, image_buf: io.BytesIO) -> None:
    image_buf.seek(0)
    _, r_id = shape.part.get_or_add_image_part(image_buf)
    blip = shape._element.find(".//" + qn("a:blip"))
    blip.set(qn("r:embed"), r_id)


def patch_tim_photo(pptx_path: Path) -> None:
    tim = Image.open(TIM_SLIDE).convert("RGB")
    prs = Presentation(str(pptx_path))
    slide = prs.slides[-1]
    pic = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    last = Image.open(io.BytesIO(pic.image.blob)).convert("RGB")
    last.paste(tim.crop(TIM_BOX), TIM_BOX[:2])
    buf = io.BytesIO()
    last.save(buf, format="JPEG", quality=88, optimize=True)
    _replace_picture(pic, buf)
    prs.save(str(pptx_path))
    print(f"  patched Tim photo on {pptx_path.name}")


def main() -> int:
    if not PDF_DIR.exists():
        print(f"PDF folder missing: {PDF_DIR}", file=sys.stderr)
        return 1
    donor = Presentation(str(DONOR))
    overlay = _overlay_elements(donor.slides[0])
    if len(overlay) != 3:
        print("donor overlay shapes missing", file=sys.stderr)
        return 1

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for item in CATALOG["templates"]:
        pdf_path = PDF_OVERRIDES.get(item["id"], PDF_DIR / item["pdf"])
        dest = OUT_DIR / f"{item['id']}.pptx"
        if not pdf_path.exists():
            print(f"missing PDF: {pdf_path}", file=sys.stderr)
            return 1
        print(item["label"])
        jpeg_dir = JPEG_SLIDES.get(item["id"])
        overlay_use = None if item["id"] in WIPE_SLIDE1 else overlay
        if jpeg_dir and jpeg_dir.exists():
            build_from_jpegs(jpeg_dir, dest, overlay_use)
        else:
            build_one(pdf_path, dest, overlay_use, template_id=item["id"])
        if item["id"] in PATCH_TIM_ON:
            patch_tim_photo(dest)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
