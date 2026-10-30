"""Fill a role deck: name + logo/photo, export PDF."""

from __future__ import annotations

import csv
import gc
import io
import json
import os
import re
import shutil
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from names import display_name

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
DEFAULT_TEMPLATE = ASSETS / "VC_Conf_MP_deck_template.pptx"
SPEC_PATH = ASSETS / "template_spec.json"
CATALOG_PATH = ASSETS / "catalog.json"
TEMPLATES_DIR = ASSETS / "templates"
FONT_FILES = ASSETS / "fonts"
USER_FONTS = Path.home() / "Library" / "Fonts"

SAFE_NAME = re.compile(r"[^A-Za-z0-9._-]+")
SAFE_FILE = re.compile(r'[<>:"/\\|?*\x00-\x1f]+')
SPILL_DIR = Path(tempfile.gettempdir()) / "vc-conf-fills"


def spill_payload(data: bytes, suffix: str) -> str:
    SPILL_DIR.mkdir(parents=True, exist_ok=True)
    fd, path = tempfile.mkstemp(suffix=suffix, prefix="vc-", dir=SPILL_DIR)
    try:
        os.write(fd, data)
    finally:
        os.close(fd)
    return path


def item_payload(item: dict) -> bytes:
    data = item.get("pdf")
    if data:
        return data
    path = item.get("pdf_path")
    if path:
        return Path(path).read_bytes()
    raise FileNotFoundError("файл сборки уже стёрт — собери ещё раз")
DEFAULT_NAME_PATTERN = "VC Conf 2027 — {name}.pdf"


def ensure_inter_installed() -> None:
    try:
        USER_FONTS.mkdir(parents=True, exist_ok=True)
        for src in FONT_FILES.glob("Inter-*.otf"):
            dest = USER_FONTS / src.name
            if not dest.exists() or dest.stat().st_size != src.stat().st_size:
                shutil.copy2(src, dest)
    except OSError:
        pass


ensure_inter_installed()


def load_spec(template_id: str | None = None) -> dict:
    data = json.loads(SPEC_PATH.read_text())
    extra = (data.get("by_id") or {}).get(template_id or "")
    spec = {k: v for k, v in data.items() if k != "by_id"}
    if not extra:
        return spec
    spec["fields"] = deepcopy(spec["fields"])
    for name, field in (extra.get("fields") or {}).items():
        base = spec["fields"].get(name) or {}
        spec["fields"][name] = {**base, **field} if isinstance(field, dict) else field
    return spec


def load_catalog() -> list[dict]:
    return json.loads(CATALOG_PATH.read_text())["templates"]


def get_template(template_id: str | None) -> dict:
    items = load_catalog()
    tid = (template_id or items[0]["id"]).strip()
    for item in items:
        if item["id"] == tid:
            path = TEMPLATES_DIR / f"{item['id']}.pptx"
            if not path.exists():
                raise FileNotFoundError(f"нет файла шаблона: {path.name}")
            return {**item, "path": path}
    raise ValueError(f"нет такого шаблона: {tid}")


def slug_filename(company: str, template_label: str | None = None) -> str:
    return render_filename(DEFAULT_NAME_PATTERN, name=company, template=template_label or "")


def render_filename(pattern: str | None, *, name: str, template: str = "") -> str:
    raw = (pattern or "").strip() or DEFAULT_NAME_PATTERN
    if not raw.lower().endswith(".pdf"):
        raw += ".pdf"
    filled = (
        raw.replace("{name}", name.strip() or "partner")
        .replace("{template}", template.strip())
        .replace("{role}", template.strip())
    )
    cleaned = SAFE_FILE.sub("", filled)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .")
    if not cleaned.lower().endswith(".pdf"):
        cleaned = (cleaned or "partner") + ".pdf"
    return cleaned or "partner.pdf"


def ascii_filename(filename: str) -> str:
    stem = Path(filename).stem
    suffix = Path(filename).suffix or ".pdf"
    slug = SAFE_NAME.sub("_", stem).strip("._") or "partner"
    return slug + suffix


def _font(weight: str, px: int) -> ImageFont.FreeTypeFont:
    file = "Inter-Bold.otf" if weight == "bold" else "Inter-Regular.otf"
    return ImageFont.truetype(str(FONT_FILES / file), px)


def _pt(pt: float, scale: float) -> int:
    return max(1, round(pt / 72 * 96 * scale))


def _round_mask(size: int, radius: int) -> Image.Image:
    mask = Image.new("L", (size, size), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, size - 1, size - 1), radius=radius, fill=255)
    return mask


def _trim_mark(im: Image.Image) -> Image.Image:
    rgba = im.convert("RGBA")
    r, g, b, a = rgba.split()
    opaque = a.point(lambda p: 255 if p > 12 else 0)
    # Drop near-white padding around a coloured mark. A white-on-transparent
    # mark has no coloured pixels, so we fall back to the alpha bbox.
    not_white = ImageChops.lighter(
        ImageChops.lighter(
            r.point(lambda p: 255 if p < 242 else 0),
            g.point(lambda p: 255 if p < 242 else 0),
        ),
        b.point(lambda p: 255 if p < 242 else 0),
    )
    mask = ImageChops.multiply(opaque, not_white)
    bbox = mask.getbbox()
    if not bbox:
        bbox = opaque.getbbox() or rgba.getbbox()
    if not bbox:
        return rgba
    x0, y0, x1, y1 = bbox
    return rgba.crop((x0, y0, x1, y1))


def _is_photo(im: Image.Image) -> bool:
    rgba = im.convert("RGBA")
    if min(rgba.size) < 160:
        return False
    sample = rgba.resize((32, 32), Image.Resampling.BOX)
    px = list(sample.getdata())
    trans = sum(1 for _r, _g, _b, a in px if a < 200)
    white = sum(1 for r, g, b, a in px if a >= 200 and r > 245 and g > 245 and b > 245)
    return trans <= 8 and white <= 10


def _flatten_white(src: Image.Image) -> Image.Image:
    rgba = src.convert("RGBA")
    bg = Image.new("RGBA", rgba.size, (255, 255, 255, 255))
    return Image.alpha_composite(bg, rgba)


def _contain_rounded(src: Image.Image, size: int, radius: int) -> Image.Image:
    mark = _trim_mark(src)
    mw, mh = mark.size
    scale = min(size / max(mw, 1), size / max(mh, 1))
    nw = max(1, round(mw * scale))
    nh = max(1, round(mh * scale))
    mark = mark.resize((nw, nh), Image.Resampling.LANCZOS)
    canvas = Image.new("RGBA", (size, size), (255, 255, 255, 255))
    canvas.paste(mark, ((size - nw) // 2, (size - nh) // 2), mark)
    canvas.putalpha(_round_mask(size, radius))
    return canvas


def _cover_rounded(src: Image.Image, size: int, radius: int) -> Image.Image:
    fitted = ImageOps.fit(_flatten_white(src), (size, size), Image.Resampling.LANCZOS)
    rgba = fitted.convert("RGBA")
    rgba.putalpha(_round_mask(size, radius))
    return rgba


def _place_mark(src: Image.Image, size: int, radius: int) -> Image.Image:
    if _is_photo(src):
        return _cover_rounded(src, size, radius)
    mark = _trim_mark(src)
    ratio = mark.width / max(mark.height, 1)
    if 0.82 <= ratio <= 1.22:
        return _cover_rounded(mark, size, radius)
    return _contain_rounded(src, size, radius)


def _wrap_text(text: str, font: ImageFont.FreeTypeFont, max_w: int) -> list[str]:
    words = text.split()
    if not words:
        return [""]
    dummy = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    lines: list[str] = []
    cur = words[0]
    for word in words[1:]:
        trial = f"{cur} {word}"
        if dummy.textlength(trial, font=font) <= max_w:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    lines.append(cur)
    return lines


def _fit_name(
    text: str,
    max_w: int,
    max_h: int,
    scale: float,
    pts: tuple[int, ...] | None = None,
) -> tuple[ImageFont.FreeTypeFont, list[str]]:
    dummy = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    spacing = round(-2 * scale)
    sizes = tuple(pts) if pts else (20, 18, 16)
    attempts = (
        (1, sizes),
        (2, sizes + (14,) if 14 not in sizes else sizes),
    )
    for max_lines, pts in attempts:
        for pt in pts:
            font = _font("bold", _pt(pt, scale))
            lines = _wrap_text(text, font, max_w)
            if len(lines) > max_lines:
                continue
            if any(dummy.textlength(line, font=font) > max_w for line in lines):
                continue
            bbox = dummy.multiline_textbbox((0, 0), "\n".join(lines), font=font, spacing=spacing)
            if (bbox[3] - bbox[1]) > max_h:
                continue
            return font, lines
    font = _font("bold", _pt(14, scale))
    return font, _wrap_text(text, font, max_w)[:2]


def _overlay(slide: Image.Image, name: str, photo: Path | io.BytesIO | None, spec: dict) -> Image.Image:
    scale = slide.width / 1920
    out = slide.convert("RGBA")
    draw = ImageDraw.Draw(out)
    logo = spec["fields"]["partner_logo"]
    x = round(logo["box_px"]["x"] * scale)
    y = round(logo["box_px"]["y"] * scale)
    w = round(logo["box_px"]["w"] * scale)
    r = round(logo["corner_radius_px"] * scale)
    ring_w = max(1, round(logo["ring"]["width_px"] * scale))

    draw.rounded_rectangle((x, y, x + w - 1, y + w - 1), radius=r, fill="#FFFFFF")
    circle = r >= max(2, w // 2 - 2)
    inset = ring_w + (max(1, round(1 * scale)) if circle else max(2, round(4 * scale)))
    ix = x + inset
    iy = y + inset
    iw = max(8, w - 2 * inset)
    ir = max(2, r - inset)
    if photo is not None:
        if hasattr(photo, "seek"):
            photo.seek(0)
        mark = _place_mark(Image.open(photo), iw, ir)
        out.paste(mark, (ix, iy), mark)
        draw = ImageDraw.Draw(out)
    draw.rounded_rectangle(
        (x, y, x + w - 1, y + w - 1),
        radius=r,
        outline=logo["ring"]["color"],
        width=ring_w,
    )

    label = spec["fields"]["partner_label"]
    if label.get("static_text"):
        lx = round(label["box_px"]["x"] * scale)
        ly = round(label["box_px"]["y"] * scale)
        label_font = _font("regular", _pt(18, scale))
        draw.text((lx, ly), label["static_text"], font=label_font, fill="#64748B")

    name_f = spec["fields"]["partner_name"]
    nx = round(name_f["box_px"]["x"] * scale)
    ny = round(name_f["box_px"]["y"] * scale)
    nw = round(name_f["box_px"]["w"] * scale)
    nh = round(name_f["box_px"]["h"] * scale)
    pts = tuple(name_f["fit_pts"]) if name_f.get("fit_pts") else None
    font, lines = _fit_name(name, nw, nh, scale, pts=pts)
    draw.multiline_text(
        (nx, ny),
        "\n".join(lines),
        font=font,
        fill=name_f["color"],
        spacing=round(-2 * scale),
    )
    return out.convert("RGB")


def _slide_picture(slide) -> Image.Image:
    pic = next(
        (s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE),
        None,
    )
    if pic is None:
        raise ValueError("в шаблоне нет картинки слайда")
    return Image.open(io.BytesIO(pic.image.blob)).convert("RGB")


def _jpeg_bytes(im: Image.Image) -> bytes:
    buf = io.BytesIO()
    im.convert("RGB").save(buf, format="JPEG", quality=88, optimize=False)
    return buf.getvalue()


def partner_preview(slide: Image.Image, spec: dict) -> bytes:
    scale = slide.width / 1920
    boxes = [
        spec["fields"]["partner_logo"]["box_px"],
        spec["fields"]["partner_name"]["box_px"],
    ]
    if spec["fields"]["partner_label"].get("static_text"):
        boxes.append(spec["fields"]["partner_label"]["box_px"])
    pad = 28
    x0 = max(0, round((min(b["x"] for b in boxes) - pad) * scale))
    y0 = max(0, round((min(b["y"] for b in boxes) - pad) * scale))
    x1 = min(slide.width, round((max(b["x"] + b["w"] for b in boxes) + pad) * scale))
    y1 = min(slide.height, round((max(b["y"] + b["h"] for b in boxes) + pad) * scale))
    crop = slide.crop((x0, y0, x1, y1)).convert("RGB")
    if crop.width < 880:
        ratio = 880 / max(crop.width, 1)
        crop = crop.resize(
            (880, max(1, round(crop.height * ratio))),
            Image.Resampling.LANCZOS,
        )
    buf = io.BytesIO()
    crop.save(buf, format="JPEG", quality=90, optimize=True)
    return buf.getvalue()


def fill_deck(
    company: str,
    logo: Path | io.BytesIO | None = None,
    template: Path | None = None,
) -> tuple[bytes, bytes]:
    import img2pdf

    tid = None
    if template is not None:
        tid = Path(template).stem
    spec = load_spec(tid)
    company = display_name(company)
    if not company:
        raise ValueError("укажи имя или название компании")

    prs = Presentation(str(template or DEFAULT_TEMPLATE))
    tmpdir = tempfile.mkdtemp(prefix="vcdeck-")
    jpeg_paths: list[str] = []
    preview = b""
    try:
        try:
            for i, slide in enumerate(prs.slides):
                im = _slide_picture(slide)
                if i == 0:
                    im = _overlay(im, company, logo, spec)
                    preview = partner_preview(im, spec)
                page = Path(tmpdir) / f"{i:03d}.jpg"
                page.write_bytes(_jpeg_bytes(im))
                jpeg_paths.append(str(page))
                im.close()
        finally:
            del prs
        return img2pdf.convert(jpeg_paths), preview
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
        gc.collect()


def _name_keys(company: str) -> list[str]:
    raw = (company or "").strip()
    if not raw:
        return []
    slug = SAFE_NAME.sub("_", raw).strip("._")
    keys = [
        raw.lower(),
        slug.lower(),
        slug.lower().replace("_", "-"),
        slug.lower().replace("_", ""),
        re.sub(r"[^a-z0-9]+", "", raw.lower()),
    ]
    seen: set[str] = set()
    out: list[str] = []
    for key in keys:
        if key and key not in seen:
            seen.add(key)
            out.append(key)
    return out


def _logo_by_company(company: str, logos_map: dict[str, bytes]) -> io.BytesIO | None:
    by_stem: dict[str, bytes] = {}
    for name, data in logos_map.items():
        stem = Path(name).stem
        by_stem[stem.lower()] = data
        slug = SAFE_NAME.sub("_", stem).strip("._").lower()
        by_stem[slug] = data
        by_stem[slug.replace("_", "-")] = data
        by_stem[slug.replace("_", "")] = data
    for key in _name_keys(company):
        if key in by_stem:
            return io.BytesIO(by_stem[key])
    return None


def _photo_from_row(value: str, logos_dir: Path | None, logos_map: dict[str, bytes]) -> io.BytesIO | None:
    if not value:
        return None
    key = Path(value.strip()).name
    if key in logos_map:
        return io.BytesIO(logos_map[key])
    if logos_dir:
        path = logos_dir / key
        if path.exists():
            return io.BytesIO(path.read_bytes())
    return None


def _row_name(lower: dict[str, str]) -> str:
    return display_name(
        lower.get("name")
        or lower.get("company")
        or lower.get("company_name")
        or lower.get("partner_name")
        or ""
    )


def fill_from_csv(
    csv_file: Path | io.TextIOBase,
    logos_dir: Path | None = None,
    logos_map: dict[str, bytes] | None = None,
    template: Path | None = None,
    template_label: str | None = None,
    name_pattern: str | None = None,
    on_progress=None,
) -> tuple[list[dict], list[dict], list[str], list[dict]]:
    from fetch_logo import fetch_logo

    logos_map = logos_map or {}
    if hasattr(csv_file, "read"):
        reader = csv.DictReader(csv_file)
        fieldnames = [h for h in (reader.fieldnames or []) if h is not None]
        rows = list(reader)
    else:
        with Path(csv_file).open(newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            fieldnames = [h for h in (reader.fieldnames or []) if h is not None]
            rows = list(reader)

    if not rows:
        raise ValueError("CSV пустой")

    total = len(rows)
    failed: list[dict] = []
    previews: list[dict] = []
    used: dict[str, int] = {}
    for i, row in enumerate(rows, start=1):
        original = {(k or ""): (v or "").strip() for k, v in row.items() if k is not None}
        lower = {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
        company = _row_name(lower)
        raw = (
            lower.get("name")
            or lower.get("company")
            or lower.get("company_name")
            or ""
        )
        label = company or raw or f"строка {i}"
        if on_progress:
            on_progress(i - 1, total, label)
        if not company:
            failed.append(
                {
                    "name": raw or f"строка {i}",
                    "url": "",
                    "reason": "нет названия",
                    "row": original,
                }
            )
            if on_progress:
                on_progress(i, total, label)
            continue
        photo_val = (
            lower.get("photo")
            or lower.get("logo")
            or lower.get("image")
            or lower.get("partner_logo")
            or ""
        )
        site = lower.get("url") or lower.get("website") or lower.get("site") or ""
        photo = _photo_from_row(photo_val, logos_dir, logos_map) if photo_val else None
        if photo is None:
            photo = _logo_by_company(company, logos_map)
        if photo is None and logos_dir:
            photo = _logo_by_company(
                company,
                {p.name: p.read_bytes() for p in logos_dir.iterdir() if p.is_file()},
            )
        if photo is None and site:
            try:
                photo = fetch_logo(site)
            except Exception:
                photo = None
        try:
            data, preview = fill_deck(company, logo=photo, template=template)
        except Exception as e:
            failed.append(
                {
                    "name": company,
                    "url": site,
                    "reason": str(e),
                    "row": original,
                }
            )
            if on_progress:
                on_progress(i, total, label)
            continue
        filename = render_filename(
            name_pattern, name=company, template=template_label or ""
        )
        if filename in used:
            used[filename] += 1
            filename = f"{Path(filename).stem} ({used[filename]}){Path(filename).suffix}"
        else:
            used[filename] = 1
        previews.append(
            {
                "name": company,
                "jpeg": preview,
                "pdf_path": spill_payload(data, ".pdf"),
                "filename": filename,
                "row": original,
                "row_index": i - 1,
            }
        )
        del data
        gc.collect()
        if on_progress:
            on_progress(i, total, company)

    if on_progress:
        on_progress(total, total, "готово")
    return failed, previews, fieldnames, rows


def pack_selected(
    previews: list[dict],
    skip: list[int],
    fieldnames: list[str],
    failed: list[dict] | None = None,
    dest: Path | None = None,
) -> bytes:
    skip_set = set(skip)
    sink: io.BufferedIOBase
    if dest is not None:
        sink = dest.open("wb")
    else:
        sink = io.BytesIO()
    try:
        with zipfile.ZipFile(sink, "w") as zf:
            for i, item in enumerate(previews):
                if i in skip_set:
                    continue
                files = item.get("files") or [item]
                for blob in files:
                    zf.writestr(
                        blob["filename"],
                        item_payload(blob),
                        compress_type=zipfile.ZIP_STORED,
                    )
            rejected = [previews[i]["row"] for i in skip if 0 <= i < len(previews)]
            for item in failed or []:
                if item.get("row"):
                    rejected.append(item["row"])
            if rejected:
                cols = fieldnames or list(rejected[0].keys())
                buf = io.StringIO()
                w = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore")
                w.writeheader()
                w.writerows(rejected)
                zf.writestr("needs_review.csv", buf.getvalue().encode("utf-8-sig"))
        if dest is not None:
            return b""
        assert isinstance(sink, io.BytesIO)
        sink.seek(0)
        return sink.getvalue()
    finally:
        if dest is not None:
            sink.close()


def links_csv(
    rows: list[dict],
    fieldnames: list[str],
    urls_by_row: dict[int, str],
    extras_by_row: dict[int, dict] | None = None,
) -> bytes:
    cols = [h for h in fieldnames if h]
    if "pdf_url" not in cols:
        cols.append("pdf_url")
    extras_by_row = extras_by_row or {}
    extra_keys: list[str] = []
    for extra in extras_by_row.values():
        for key in extra:
            if key not in cols and key not in extra_keys:
                extra_keys.append(key)
    cols.extend(extra_keys)
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore")
    w.writeheader()
    for i, row in enumerate(rows):
        out = {k: (row.get(k) or "") for k in fieldnames if k}
        out["pdf_url"] = urls_by_row.get(i, "")
        out.update(extras_by_row.get(i) or {})
        w.writerow(out)
    return buf.getvalue().encode("utf-8-sig")


def unpack_logo_zip(data: bytes) -> dict[str, bytes]:
    out: dict[str, bytes] = {}
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            name = Path(info.filename).name
            if name.startswith(".") or not name:
                continue
            out[name] = zf.read(info)
    return out
